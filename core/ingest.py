"""
ingest.py — Document extraction and chunking.

Formats: PDF, Word, PowerPoint, Outlook messages, plain email, Markdown, text.

The two older Microsoft formats, .doc and .ppt, are deliberately absent. No
pure Python library reads them reliably; support would mean depending on
LibreOffice or antiword being installed, which most people running this will
not have, and a silent failure is worse than a clear refusal. Opening such a
file and saving it in the current format takes a moment and is something the
person can do themselves.
"""

import email
from email import policy
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".pptx", ".txt", ".md", ".eml", ".msg",
}

# Shown to a person choosing files, so they know before they try.
FORMAT_NOTE = "PDF, Word (.docx), PowerPoint (.pptx), email (.eml, .msg), text, Markdown"


def extract_pdf(path: Path) -> str | None:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        parts = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"\n[Page {i}]\n{text}")
        return "".join(parts) or None
    except Exception:
        return None


def extract_docx(path: Path) -> str | None:
    """
    Word document text, including tables.

    Tables carry a lot of the substance in consulting documents, so their cells
    are pulled out rather than skipped.
    """
    try:
        from docx import Document
        doc = Document(str(path))
        parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                style = (para.style.name or "").lower()
                if style.startswith("heading"):
                    parts.append(f"\n## {para.text}")
                else:
                    parts.append(para.text)

        for t_index, table in enumerate(doc.tables, 1):
            rows = []
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"\n[Table {t_index}]\n" + "\n".join(rows))

        return "\n".join(parts) or None
    except Exception:
        return None


def extract_pptx(path: Path) -> str | None:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            slide_text = ""
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text += shape.text + "\n"
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        slide_text += " | ".join(c.text.strip() for c in row.cells) + "\n"
            if slide_text.strip():
                parts.append(f"\n[Slide {i}]\n{slide_text}")
        return "".join(parts) or None
    except Exception:
        return None


def extract_eml(path: Path) -> str | None:
    try:
        with open(path, "rb") as f:
            msg = email.message_from_binary_file(f, policy=policy.default)
        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    try:
                        body += part.get_content() + "\n"
                    except Exception:
                        pass
        else:
            try:
                body = msg.get_content()
            except Exception:
                body = ""
        subject = str(msg.get("Subject", ""))
        if not body.strip() and not subject.strip():
            return None
        return (
            f"From: {msg.get('From', '')}\n"
            f"To: {msg.get('To', '')}\n"
            f"Date: {msg.get('Date', '')}\n"
            f"Subject: {subject}\n\n{body}"
        )
    except Exception:
        return None


def extract_msg(path: Path) -> str | None:
    """
    Outlook message text.

    Kept in the same shape as .eml output so a mailbox exported either way
    reads consistently once indexed.
    """
    try:
        import extract_msg
        with extract_msg.Message(str(path)) as msg:
            body = msg.body or ""
            subject = msg.subject or ""
            if not body.strip() and not subject.strip():
                return None
            return (
                f"From: {msg.sender or ''}\n"
                f"To: {msg.to or ''}\n"
                f"Date: {msg.date or ''}\n"
                f"Subject: {subject}\n\n{body}"
            )
    except Exception:
        return None


def extract_text(path: Path) -> str | None:
    try:
        return path.read_text(encoding="utf-8", errors="replace") or None
    except Exception:
        return None


_EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".eml": extract_eml,
    ".msg": extract_msg,
    ".txt": extract_text,
    ".md": extract_text,
}


def extract(path: Path) -> str | None:
    """Text from a supported file, or None if it cannot be read."""
    handler = _EXTRACTORS.get(path.suffix.lower())
    return handler(path) if handler else None


def missing_reader(path: Path) -> str | None:
    """
    The package needed for this file type, if it is not installed.

    Extraction returns None on any failure, which cannot distinguish a corrupt
    file from a missing library. This says which it is, so the fix is a pip
    install rather than a guess.
    """
    needed = {
        ".pdf": ("pypdf", "pypdf"),
        ".docx": ("docx", "python-docx"),
        ".pptx": ("pptx", "python-pptx"),
        ".msg": ("extract_msg", "extract-msg"),
    }.get(path.suffix.lower())
    if not needed:
        return None
    module, package = needed
    try:
        __import__(module)
        return None
    except ImportError:
        return package


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 100) -> list[str]:
    """Split text into overlapping chunks. Overlap must be smaller than the size."""
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")
    text = text.encode("utf-8", errors="replace").decode("utf-8")
    if len(text) <= chunk_size:
        return [text] if text.strip() else []
    chunks, start, step = [], 0, chunk_size - overlap
    while start < len(text):
        chunk = text[start:start + chunk_size]
        if chunk.strip():
            chunks.append(chunk)
        start += step
    return chunks


def scan_directory(directory: Path, recursive: bool = True) -> list[Path]:
    """Every supported file in a directory."""
    if not directory.exists():
        return []
    pattern = "**/*" if recursive else "*"
    return sorted(
        p for p in directory.glob(pattern)
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )
